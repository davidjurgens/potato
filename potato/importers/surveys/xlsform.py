"""
XLSForm spreadsheets (ODK, KoboToolbox, Ona, SurveyCTO) into Potato survey pages.

An XLSForm is a workbook with a ``survey`` sheet (columns ``type``,
``name``, ``label``, ``relevant``, ``required``, ``appearance``, ...), a
``choices`` sheet (``list_name``, ``name``, ``label``) and an optional
``settings`` sheet (``form_title``, ``default_language``). Labels may be
per language (``label::English (en)``); the importer uses
``options["language"]``, else ``default_language``, else the plain
``label`` column, else the first language.

Pages. ODK shows one question per screen unless a group has the
``field-list`` (or ``table-list``) appearance, which puts the whole group
on one screen. Potato pages hold several questions, so the importer puts
each field-list group on its own page and every run of questions between
them on one page. Groups without field-list do not start a page; their
label is shown as a heading.

Mapped:

* ``select_one`` to radio (``minimal`` / ``autocomplete`` appearance to a
  dropdown, ``likert`` and ``horizontal`` / ``columns`` / ``compact`` to a
  horizontal radio), ``select_multiple`` to multiselect, ``rank`` to
  ranking, ``integer`` / ``decimal`` to number, ``range`` to slider
  (``parameters`` ``start`` / ``end`` / ``step``), ``text`` to text
  (``multiline`` to a textarea), ``note`` to a display block,
  ``acknowledge`` to a one-option ("OK") radio; ``or_other`` lists get an
  "Other" choice and a text box; the legacy spellings (``select one from``,
  ``select all that apply``, ``string``, ``int``, ``trigger``).
* ``required`` (yes / true / true() / 1), ``parameters: randomize=true``.
* ``relevant`` on questions and groups (a group's condition ANDed onto
  everything inside it; on a field-list page it becomes the page's
  ``show_when``) through the shared expression translator: comparisons,
  ``and`` / ``or``, ``not()``, ``selected()``, ``count-selected()``,
  ``string-length()``, ``regex()``, and ``${q} = ''`` as "empty".

Reported, not mapped: date, time, dateTime, geopoint / geotrace / geoshape,
image, audio, video, file and barcode questions; ``calculate`` and
``hidden`` fields (and logic that refers to them); ``select_*_from_file``
and external choices; repeats (their questions are imported once);
``constraint``, ``choice_filter``, ``default``, ``read_only``,
``calculation``, conditional ``required``; piped text (``${q}`` in a label);
and logic that calls other functions (``jr:choice-name``, ``if``,
``coalesce``, ``concat``, ``today``, ``now``, ``position``, ...). Metadata
fields (``start``, ``end``, ``today``, ``deviceid``, ...) are skipped
without a report line: they record the session, they are not questions.

``.xls`` workbooks are read when the ``xlrd`` package is installed;
otherwise the importer asks for the file to be saved as ``.xlsx``.
"""

from __future__ import annotations

import datetime as _dt
import re
from pathlib import Path
from typing import Any, Dict, List, Optional

from .base import (
    NameAllocator, SurveyImportResult, SurveyPage, Untranslatable, and_rules,
    strip_html, unique_labels)
from .expr import XLSFORM, VarInfo, compare, parse, to_item, to_rule

_HEADER_ALIASES = {
    "list name": "list_name", "list_name": "list_name",
    "relevance": "relevant", "jr:constraintmsg": "constraint_message",
    "constraint_msg": "constraint_message", "requiredmsg": "required_message",
    "jr:requiredmsg": "required_message", "readonly": "read_only",
    "choice_filter": "choice_filter", "calculate": "calculation",
    "image": "media_image", "audio": "media_audio", "video": "media_video",
    "repeat_count": "repeat_count", "jr:count": "repeat_count",
    "caption": "label",
}

_METADATA = {
    "start", "end", "today", "deviceid", "subscriberid", "simserial",
    "phonenumber", "username", "email", "audit", "imei", "start-geopoint",
    "background-geopoint", "uri:deviceid", "uri:subscriberid", "uri:simserial",
    "uri:phonenumber", "uri:email", "uri:username", "instanceid",
}

_UNSUPPORTED = {
    "date": "date", "time": "time", "datetime": "date and time",
    "geopoint": "location", "geotrace": "line on a map", "geoshape": "area on a map",
    "image": "photo", "audio": "audio recording", "video": "video recording",
    "file": "file upload", "barcode": "barcode scan", "background-audio": "background audio",
    "xml-external": "external XML data", "csv-external": "external CSV data",
    "osm": "OpenStreetMap", "begin_loop": "loop", "end_loop": "loop",
    "select_one_from_file": "choices from a file", "select_multiple_from_file": "choices from a file",
    "select_one_external": "external choices",
}

_TYPE_ALIASES = {
    "string": "text", "int": "integer", "trigger": "acknowledge",
}

_LEGACY_PREFIXES = [
    ("select all that apply from ", "select_multiple "),
    ("select all that apply ", "select_multiple "),
    ("select multiple from ", "select_multiple "),
    ("select one from ", "select_one "),
    ("select one ", "select_one "),
    ("select1 ", "select_one "),
    ("begin group", "begin_group"), ("end group", "end_group"),
    ("begin repeat", "begin_repeat"), ("end repeat", "end_repeat"),
    ("begin loop", "begin_loop"), ("end loop", "end_loop"),
]

_YES = {"yes", "true", "true()", "1", "y"}
_NO = {"", "no", "false", "false()", "0", "n"}
_PIPE = re.compile(r"\$\{[^{}]+\}")


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (_dt.datetime, _dt.date, _dt.time)):
        return value.isoformat()
    return str(value).strip()


def _norm_header(h: Any) -> str:
    h = _cell(h)
    for prefix in ("bind::", "control::", "body::", "media::"):
        if h.lower().startswith(prefix):
            h = h[len(prefix):]
    base, sep, lang = h.partition("::")
    base = base.strip().lower()
    base = _HEADER_ALIASES.get(base, base.replace(" ", "_"))
    return f"{base}::{lang.strip()}" if sep else base


def _quotes(text: str) -> str:
    """Spreadsheet autocorrect turns ' into curly quotes inside expressions."""
    return (text.replace("‘", "'").replace("’", "'")
            .replace("“", '"').replace("”", '"'))


def _parameters(text: str) -> Dict[str, str]:
    out = {}
    for part in re.split(r"[\s,;]+", text or ""):
        if "=" in part:
            k, _, v = part.partition("=")
            out[k.strip().lower()] = v.strip()
    return out


def _num(value: Any, default=None):
    try:
        v = float(str(value).strip())
    except (TypeError, ValueError):
        return default
    return int(v) if v.is_integer() else v


def _read_workbook(path: Path) -> Dict[str, List[List[Any]]]:
    """{sheet name (lower-cased): rows}, from .xlsx (openpyxl) or .xls (xlrd)."""
    sheets: Dict[str, List[List[Any]]] = {}
    if path.suffix.lower() == ".xls":
        try:
            import xlrd  # noqa: F401
        except ImportError:
            raise ValueError(f"{path.name} is a legacy .xls workbook; open it in a spreadsheet "
                             "program and save it as .xlsx, then import that file")
        try:
            book = xlrd.open_workbook(str(path))
        except Exception as e:  # noqa: BLE001 - xlrd raises its own error types
            raise ValueError(f"{path.name}: not a readable .xls workbook ({e})")
        for sheet in book.sheets():
            rows = []
            for r in range(sheet.nrows):
                row = []
                for c in range(sheet.ncols):
                    cell = sheet.cell(r, c)
                    v = cell.value
                    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        v = bool(v)
                    row.append(v)
                rows.append(row)
            sheets.setdefault(sheet.name.strip().lower(), rows)
        return sheets
    import openpyxl
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            sheets.setdefault(ws.title.strip().lower(),
                              [list(r) for r in ws.iter_rows(values_only=True)])
    finally:
        wb.close()
    return sheets


def _records(rows: List[List[Any]]):
    """(row number, {normalized header: cell text}) for each non-empty row."""
    header = None
    for i, row in enumerate(rows, start=1):
        cells = [_cell(c) for c in row]
        if not any(cells):
            continue
        if header is None:
            header = [_norm_header(c) for c in row]
            continue
        rec: Dict[str, str] = {}
        for key, val in zip(header, cells):
            if key and val and key not in rec:
                rec[key] = val
        if rec:
            yield i, rec


def _has_survey_header(rows: List[List[Any]]) -> bool:
    for row in rows:
        cells = [_norm_header(c) for c in row]
        if any(cells):
            return "type" in cells and "name" in cells
    return False


class XLSFormImporter:
    format_name = "xlsform"
    description = "XLSForm spreadsheet (ODK, KoboToolbox, Ona, SurveyCTO)"
    file_extensions = [".xlsx", ".xls"]

    def detect_path(self, path: Path) -> bool:
        path = Path(path)
        if not path.is_file():
            return False
        suffix = path.suffix.lower()
        if suffix == ".xls":
            with path.open("rb") as f:
                return f.read(8) == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
        if suffix != ".xlsx":
            return False
        try:
            import openpyxl
            wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        except Exception:  # noqa: BLE001 - any unreadable workbook is "not ours"
            return False
        try:
            ws = next((w for w in wb.worksheets if w.title.strip().lower() == "survey"), None)
            if ws is None:
                return False
            # Only the header row is needed; read up to the first non-empty one.
            head = []
            for row in ws.iter_rows(values_only=True, max_row=20):
                head.append(list(row))
                if any(_cell(c) for c in row):
                    break
            return _has_survey_header(head)
        except Exception:  # noqa: BLE001
            return False
        finally:
            wb.close()

    def parse_path(self, path: Path, options: Optional[dict] = None) -> SurveyImportResult:
        path = Path(path)
        sheets = _read_workbook(path)
        return self.parse(sheets, options or {}, path.stem)

    # ------------------------------------------------------------------

    def parse(self, sheets: Dict[str, List[List[Any]]], options: dict,
              stem: str = "form") -> SurveyImportResult:
        if "survey" not in sheets or not _has_survey_header(sheets["survey"]):
            raise ValueError("not an XLSForm: expected a 'survey' sheet with 'type' and 'name' columns")
        survey = list(_records(sheets["survey"]))
        choices_rows = list(_records(sheets.get("choices", [])))
        settings = next((rec for _i, rec in _records(sheets.get("settings", []))), {})

        result = SurveyImportResult(source_format="XLSForm",
                                    title=settings.get("form_title", "") or "")
        self.result = result
        self.lang = self._language(survey, options, settings)
        if self.lang:
            result.warnings.append(f"labels taken from the '{self.lang}' language columns")

        self.choices: Dict[str, List[Dict[str, str]]] = {}
        media_only = set()
        for _i, rec in choices_rows:
            list_name = rec.get("list_name", "")
            if not list_name or "name" not in rec:
                continue
            label = self._label(rec)
            if not label and any(k.startswith("media_") for k in rec):
                media_only.add(list_name)
            self.choices.setdefault(list_name, []).append(
                {**rec, "label": strip_html(label) or rec["name"]})
        for list_name in sorted(media_only):
            result.skip(f"choice list '{list_name}'", "image/audio choices", "imported with their names as text")

        self.names = NameAllocator()
        self.page_names = NameAllocator(reserved={"annotation", "order", "counterbalance"})
        self.vars: Dict[str, VarInfo] = {}
        self.blocked: Dict[str, str] = {}
        self.stack: List[Dict[str, Any]] = []
        self.run_page: Optional[SurveyPage] = None
        self.coverage: Dict[int, set] = {}
        self.pending: List[tuple] = []   # (target, kind, [(expr, where)])
        self.run_count = 0

        for rownum, rec in survey:
            self._row(rownum, rec)
        if self.stack:
            result.warnings.append(f"{len(self.stack)} group(s) or repeat(s) never closed")

        for target, kind, exprs in self.pending:
            rule = None
            for expression, where in exprs:
                rule = and_rules(rule, self._relevant(expression, where))
            if rule is None:
                continue
            if kind == "page":
                target.show_when = and_rules(target.show_when, rule)
            else:
                target["display_logic"] = and_rules(target.get("display_logic"), rule)
        for key in ("instance_name", "public_key", "submission_url"):
            if settings.get(key):
                result.skip("settings", key, settings[key][:200])
        return result

    def _language(self, survey, options, settings) -> Optional[str]:
        langs: List[str] = []
        plain = False
        for _i, rec in survey:
            for key in rec:
                if key.startswith("label::"):
                    lang = key[7:]
                    if lang not in langs:
                        langs.append(lang)
                elif key == "label":
                    plain = True
        def find(want):
            # "French (fr)" is found by itself, by "French" or by "fr".
            w = str(want).strip().lower()
            for lang in langs:
                name, _, code = lang.partition("(")
                if w in (lang.lower(), name.strip().lower(), code.rstrip(")").strip().lower()):
                    return lang
            return None

        asked = options.get("language")
        if asked and find(asked):
            return find(asked)
        default = settings.get("default_language")
        chosen = find(default) if default else None
        if chosen is None and not plain and langs:
            chosen = langs[0]
        if asked:
            self.result.warnings.append(
                f"--survey-language '{asked}' is not one of this form's languages "
                f"({', '.join(langs) or 'it has none'}); using "
                + (f"'{chosen}'" if chosen else "the plain label column"))
        return chosen

    def _label(self, rec: Dict[str, str], key: str = "label") -> str:
        if self.lang and rec.get(f"{key}::{self.lang}"):
            return rec[f"{key}::{self.lang}"]
        if rec.get(key):
            return rec[key]
        for k, v in rec.items():
            if k.startswith(f"{key}::") and v:
                return v
        return ""

    # ------------------------------------------------------------------
    # Structure

    def _page(self) -> SurveyPage:
        """The page the next question goes on: the outermost open field-list
        group's page, else the current run of top-level questions."""
        top = next((f for f in self.stack if f["fieldlist"]), None)
        if top is not None:
            return self._page_of(top)
        if self.run_page is None:
            self.run_count += 1
            self.run_page = self._new_page(f"page{self.run_count}", "", [])
        return self.run_page

    def _page_of(self, frame) -> SurveyPage:
        if frame["page"] is None:
            frame["page"] = self._new_page(frame["name"], frame["label"],
                                           self.stack[:self.stack.index(frame) + 1])
        return frame["page"]

    def _new_page(self, name: str, title: str, covered_frames: List[Dict[str, Any]]) -> SurveyPage:
        page = SurveyPage(name=self.page_names.take(name, "page"), title=title)
        self.result.pages.append(page)
        self.coverage[id(page)] = {id(f) for f in covered_frames}
        exprs = [(f["relevant"], f["where"]) for f in covered_frames if f["relevant"]]
        if exprs:
            self.pending.append((page, "page", exprs))
        return page

    def _add(self, scheme: Dict[str, Any], relevant: str, where: str) -> None:
        page = self._page()
        page.questions.append(scheme)
        covered = self.coverage.get(id(page), set())
        exprs = [(f["relevant"], f["where"]) for f in self.stack
                 if f["relevant"] and id(f) not in covered]
        if relevant:
            exprs.append((relevant, where))
        if exprs:
            self.pending.append((scheme, "question", exprs))

    # ------------------------------------------------------------------
    # Rows

    def _row(self, rownum: int, rec: Dict[str, str]) -> None:
        raw_type = " ".join(rec.get("type", "").split())
        low = raw_type.lower()
        for old, new in _LEGACY_PREFIXES:
            if low.startswith(old):
                low = new + low[len(old):]
                raw_type = new + raw_type[len(old):]
                break
        parts = raw_type.split()
        base = parts[0].lower() if parts else ""
        base = _TYPE_ALIASES.get(base, base)
        src = rec.get("name", "")
        where = f"row {rownum}" + (f" '{src}'" if src else "")
        relevant = _quotes(rec.get("relevant", ""))

        if not base:
            if src or self._label(rec):
                self.result.skip(where, "row without a type")
            return
        if base in ("begin_group", "begin_repeat"):
            appearance = rec.get("appearance", "").lower().split()
            frame = {"name": src or ("repeat" if base == "begin_repeat" else "group"),
                     "label": strip_html(self._label(rec)), "relevant": relevant,
                     "where": f"{where} ({'repeat' if base == 'begin_repeat' else 'group'})",
                     "fieldlist": bool({"field-list", "table-list"} & set(appearance)),
                     "page": None, "repeat": base == "begin_repeat"}
            if base == "begin_repeat":
                detail = "its questions are imported once"
                if rec.get("repeat_count"):
                    detail += f"; repeat_count {rec['repeat_count']}"
                self.result.skip(where, "repeat", detail)
            self.stack.append(frame)
            if frame["label"]:
                self._add({"name": self.names.take(f"{frame['name']}_title", "group"),
                           "annotation_type": "pure_display", "description": frame["label"]},
                          "", where)
            return
        if base in ("end_group", "end_repeat"):
            if self.stack:
                frame = self.stack.pop()
                if frame["fieldlist"] and not any(f["fieldlist"] for f in self.stack):
                    self.run_page = None
            return
        if base in _METADATA:
            self.blocked[src] = f"{src} is a metadata field ({base})"
            return
        if base in ("calculate", "hidden"):
            self.blocked[src] = f"{src} is a {base} field"
            detail = rec.get("calculation", "") or rec.get("default", "")
            self.result.skip(where, f"{base} field", detail[:200])
            return
        if base in _UNSUPPORTED:
            self.blocked[src] = f"{src} is a {_UNSUPPORTED[base]} question"
            self.result.skip(where, f"question type '{raw_type}'", strip_html(self._label(rec))[:200])
            return

        scheme = self._scheme(base, parts, rec, src, where)
        if scheme is None:
            return
        label = self._label(rec)
        hint = strip_html(self._label(rec, "hint"))
        text = strip_html(label) or src
        scheme["description"] = f"{text} ({hint})" if hint else text
        if _PIPE.search(label or ""):
            self.result.skip(where, "piped text", label[:200])
        required = rec.get("required", "").strip().lower()
        if required in _YES and scheme["annotation_type"] != "pure_display":
            scheme["label_requirement"] = {"required": True}
        elif required not in _NO and required not in _YES:
            self.result.skip(where, "conditional required", rec["required"])
        for key, what in (("constraint", "constraint"), ("choice_filter", "choice filter"),
                          ("default", "default value"), ("calculation", "calculation"),
                          ("trigger", "trigger")):
            if rec.get(key):
                detail = rec[key]
                msg = self._label(rec, "constraint_message") if key == "constraint" else ""
                self.result.skip(where, what, f"{detail} ({msg})" if msg else detail)
        if rec.get("read_only", "").lower() in _YES:
            self.result.skip(where, "read-only question")
        self.result.name_map[src] = scheme["name"]
        self._add(scheme, relevant, where)

    def _scheme(self, base: str, parts: List[str], rec: Dict[str, str], src: str,
                where: str) -> Optional[Dict[str, Any]]:
        appearance = rec.get("appearance", "").lower().split()
        params = _parameters(rec.get("parameters", ""))
        randomize = params.get("randomize", "").lower() in _YES

        if base in ("select_one", "select_multiple", "rank"):
            if len(parts) < 2:
                self.result.skip(where, f"{base} without a choice list")
                return None
            list_name = parts[1]
            or_other = len(parts) > 2 and parts[2].lower() == "or_other"
            if list_name not in self.choices:
                self.result.skip(where, f"choice list '{list_name}' not found in the choices sheet")
                self.blocked[src] = f"{src} has no choices"
                return None
            items = self.choices[list_name]
            labels = unique_labels([c["label"] for c in items] + (["Other"] if or_other else []))
            values = {c["name"]: lab for c, lab in zip(items, labels)}
            if or_other:
                values["other"] = labels[-1]
            name = self.names.take(src)
            scheme: Dict[str, Any] = {"name": name, "labels": labels}
            if or_other:
                scheme["has_free_response"] = {"instruction": "Please specify"}
            if any(a.startswith("search(") for a in appearance) or rec.get("appearance", "").startswith("search("):
                self.result.skip(where, "search() appearance", rec["appearance"])
            if base == "rank":
                scheme["annotation_type"] = "ranking"
                scheme.pop("has_free_response", None)
                self.blocked[src] = f"{src} is a ranking; logic on ranks is not translated"
                if randomize:
                    self.result.skip(where, "randomize", "ranking options keep their order")
                return scheme
            if base == "select_multiple":
                scheme["annotation_type"] = "multiselect"
                self.vars[src] = VarInfo(schema=name, kind="multi", values=values)
            else:
                horizontal = any(a == "likert" or a.startswith(("horizontal", "columns", "compact"))
                                 or a in ("quickcompact", "label", "list-nolabel") for a in appearance)
                dropdown = any(a in ("minimal", "autocomplete", "autocomplete_chars")
                               for a in appearance)
                scheme["annotation_type"] = "select" if dropdown else "radio"
                if dropdown and or_other:
                    scheme.pop("has_free_response")
                    self.result.skip(where, "free-text 'or_other' on a dropdown",
                                     "imported as an 'Other' option without the text box")
                if horizontal and not dropdown:
                    scheme["horizontal"] = True
                self.vars[src] = VarInfo(schema=name, kind="single", values=values)
            if randomize:
                scheme["randomize_order"] = True
            return scheme
        if base in ("integer", "decimal"):
            name = self.names.take(src)
            self.vars[src] = VarInfo(schema=name, kind="number")
            return {"name": name, "annotation_type": "number", "step": 1 if base == "integer" else "any"}
        if base == "range":
            start = _num(params.get("start"), 1)
            end = _num(params.get("end"), 10)
            step = _num(params.get("step"), 1)
            if start > end:
                start, end = end, start
                self.result.warnings.append(f"{where}: descending range imported ascending")
            name = self.names.take(src)
            self.vars[src] = VarInfo(schema=name, kind="number")
            if rec.get("appearance"):
                self.result.warnings.append(f"{where}: range appearance '{rec['appearance']}' "
                                            "imported as a plain slider")
            return {"name": name, "annotation_type": "slider", "min_value": start,
                    "max_value": end, "starting_value": start, "step": step}
        if base == "text":
            name = self.names.take(src)
            self.vars[src] = VarInfo(schema=name, kind="text")
            scheme = {"name": name, "annotation_type": "text"}
            if "multiline" in appearance:
                scheme["textarea"] = {"on": True, "rows": _num(rec.get("rows"), 4) or 4}
            if "numbers" in appearance:
                self.result.warnings.append(f"{where}: 'numbers' appearance imported as a plain text box")
            return scheme
        if base == "note":
            name = self.names.take(src)
            self.blocked[src] = f"{src} is a note, not a question"
            return {"name": name, "annotation_type": "pure_display"}
        if base == "acknowledge":
            name = self.names.take(src)
            self.vars[src] = VarInfo(schema=name, kind="single", values={"OK": "OK"})
            self.result.warnings.append(f"{where}: acknowledge imported as a radio with the single option 'OK'")
            return {"name": name, "annotation_type": "radio", "labels": ["OK"]}
        self.blocked[src] = f"{src} has the unknown type '{base}'"
        self.result.skip(where, f"question type '{' '.join(parts)}'", strip_html(self._label(rec))[:200])
        return None

    # ------------------------------------------------------------------
    # Logic

    def _resolve(self, ref: str) -> VarInfo:
        ref = ref.strip()
        if ref == ".":
            raise Untranslatable("'.' (the current question) in a relevant expression")
        if ref in self.vars:
            return self.vars[ref]
        if ref in self.blocked:
            raise Untranslatable(self.blocked[ref])
        raise Untranslatable(f"unknown question {ref!r}")

    def _call(self, name: str, args: list, resolve, cmp) -> Optional[Dict[str, Any]]:
        low = name.lower()
        if low == "not" and cmp is None and len(args) == 1:
            return {"not": to_item(args[0], resolve, self._call)}
        if low == "selected" and cmp is None and len(args) == 2 \
                and args[0][0] == "var" and args[1][0] == "lit":
            info = resolve(args[0][1])
            if info.kind == "multi":
                return compare(info, "contains", args[1][1])
            if info.kind == "single":
                return compare(info, "=", args[1][1])
            raise Untranslatable(f"selected() on {info.schema}, which is not a choice question")
        if low == "count-selected" and cmp is not None and len(args) == 1 and args[0][0] == "var":
            op, right = cmp
            if right[0] != "lit":
                return None
            info = resolve(args[0][1])
            return compare(VarInfo(schema=info.schema, kind="count"), op, right[1])
        if low == "string-length" and cmp is not None and len(args) == 1 and args[0][0] == "var":
            op, right = cmp
            n = _num(right[1]) if right[0] == "lit" else None
            if not isinstance(n, int):
                return None
            info = resolve(args[0][1])
            leaf = {"schema": info.schema}
            if op == ">":
                return {**leaf, "operator": "length_gt", "value": n}
            if op == ">=":
                return {**leaf, "operator": "length_gt", "value": n - 1}
            if op == "<":
                return {**leaf, "operator": "length_lt", "value": n}
            if op == "<=":
                return {**leaf, "operator": "length_lt", "value": n + 1}
            if op == "=":
                return {**leaf, "operator": "length_in_range", "value": [n, n]}
            if op == "!=":
                return {"not": {**leaf, "operator": "length_in_range", "value": [n, n]}}
            return None
        if low == "regex" and cmp is None and len(args) == 2 \
                and args[0][0] == "var" and args[1][0] == "lit":
            info = resolve(args[0][1])
            return {"schema": info.schema, "operator": "matches", "value": str(args[1][1])}
        return None

    def _relevant(self, expression: str, where: str):
        try:
            # XPath strings have no escapes: keep regex backslashes (\\.) intact
            # through the shared tokenizer, which unescapes them.
            tree = parse(expression, XLSFORM)
            if tree == ("lit", True):
                return None
            if tree == ("lit", False):
                raise Untranslatable("never shown (false())")
            return to_rule(tree, self._resolve, self._call)
        except Untranslatable as e:
            self.result.skip(where, "relevant", f"{expression} ({e})")
            return None
